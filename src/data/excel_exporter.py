import os
import math
from datetime import datetime
from typing import List, Optional, Any, Dict, Tuple

import numpy as np

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from ..analysis.annotations import Annotation, AnnotationManager
from ..utils.math_utils import adaptive_log_y_range, _py_scalar
from ..config import RGB_Y_MIN, RGB_Y_MAX                                                           

def _hex_to_fill(hex_color: str) -> PatternFill:
    hex_clean = str(hex_color or '').lstrip('#').upper()
    if len(hex_clean) != 6:
        hex_clean = 'FF4444'
    return PatternFill(
        start_color=hex_clean, end_color=hex_clean, fill_type="solid"
    )


def _marker_table_headers(has_log: bool = True) -> List[str]:
    headers = [
        '#', 'На графике (с)', 'На графике (мин)', 'Создано',
        'R', 'G', 'B',
    ]
    if has_log:
        headers.extend(['Log10(B/R)', 'Log10(B/G)'])
    headers.extend(['Описание', 'Цвет'])
    return headers


def _marker_table_widths(has_log: bool = True) -> List[int]:
    widths = [5, 12, 14, 20, 8, 8, 8]
    if has_log:
        widths.extend([12, 12])
    widths.extend([30, 8])
    return widths


def _nearest_saved_data_index(saved_data, time_ms):
    """Find the nearest row in an ordered session without allocating an O(N) array."""
    if not saved_data:
        return None
    try:
        target = float(time_ms)
        low, high = 0, len(saved_data)
        while low < high:
            middle = (low + high) // 2
            if float(saved_data[middle][0]) < target:
                low = middle + 1
            else:
                high = middle
        candidates = [idx for idx in (low - 1, low) if 0 <= idx < len(saved_data)]
        # np.argmin used to prefer the first row on equal distances.
        return min(candidates, key=lambda idx: (abs(float(saved_data[idx][0]) - target), idx))
    except (TypeError, ValueError, IndexError, OverflowError):
        times = np.asarray([float(row[0]) for row in saved_data], dtype=np.float64)
        return int(np.argmin(np.abs(times - float(time_ms))))


def _annotation_values_at_time(saved_data: List, time_ms: float, has_log: bool) -> Dict:
    if not saved_data:
        return {'r': None, 'g': None, 'b': None, 'log_br': None, 'log_bg': None}
    idx = _nearest_saved_data_index(saved_data, time_ms)
    row = saved_data[idx]
    values = {
        'r': _py_scalar(row[1]),
        'g': _py_scalar(row[2]),
        'b': _py_scalar(row[3]),
        'log_br': None,
        'log_bg': None,
    }
    if has_log and len(row) > 4:
        values['log_br'] = _py_scalar(row[4])
    if has_log and len(row) > 5:
        values['log_bg'] = _py_scalar(row[5])
    return values


def _normalize_annotations(annotations) -> List[Annotation]:
    if not annotations:
        return []
    result = []
    for item in annotations:
        if isinstance(item, Annotation):
            result.append(item)
        elif isinstance(item, dict):
            result.append(Annotation.from_dict(item))
    return result


def _filter_annotations(annotations, is_auto: Optional[bool] = None) -> List[Annotation]:
    annotations = _normalize_annotations(annotations)
    if is_auto is None:
        return sorted(annotations, key=lambda a: a.time_ms)
    return sorted([ann for ann in annotations if ann.is_auto == is_auto], key=lambda a: a.time_ms)


def _fill_annotation_sheet_headers(ws, has_log: bool = True):
    headers = _marker_table_headers(has_log)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for col, width in enumerate(_marker_table_widths(has_log), 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _fill_auto_annotation_sheet_headers(ws):
    headers = [
        '#', 'На графике (с)', 'Минуты', 'R', 'G', 'B',
        'RGB_sum_slope_30s', 'Transition_score', 'Описание', 'Цвет'
    ]
    widths = [5, 12, 14, 8, 8, 8, 18, 16, 30, 12]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_auto_annotations_to_sheet(ws, annotations):
    annotations = _filter_annotations(annotations, is_auto=True)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_idx, ann in enumerate(annotations, 2):
        seconds = None
        minutes = None
        try:
            seconds = float(ann.time_ms) / 1000.0
        except Exception:
            seconds = None
        try:
            minutes = float(ann.time_ms) / 60000.0
        except Exception:
            minutes = None

        row_values = [
            row_idx - 1,
            _py_scalar(round(seconds, 3) if seconds is not None else None),
            _py_scalar(round(minutes, 3) if minutes is not None else None),
            _py_scalar(ann.r),
            _py_scalar(ann.g),
            _py_scalar(ann.b),
            _py_scalar(ann.rgb_sum_slope_30s),
            _py_scalar(ann.transition_score),
            ann.description,
        ]
        for col, value in enumerate(row_values, 1):
            ws.cell(row=row_idx, column=col, value=value)

        color_cell = ws.cell(row=row_idx, column=len(row_values) + 1, value="")
        color_cell.fill = _hex_to_fill(ann.color)


def _write_annotations_to_sheet(ws, annotations, saved_data: Optional[List] = None, has_log: bool = True):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    sorted_anns = _filter_annotations(annotations, is_auto=False)
    headers = _marker_table_headers(has_log)
    color_col = len(headers)
    row_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for row_idx, ann in enumerate(sorted_anns, 2):
        seconds = None
        minutes = None
        try:
            seconds = float(ann.time_ms) / 1000.0
        except Exception:
            seconds = None
        try:
            minutes = float(ann.time_ms) / 60000.0
        except Exception:
            minutes = None
        vals = _annotation_values_at_time(saved_data, ann.time_ms, has_log)
        row_values = [
            row_idx - 1,
            _py_scalar(round(seconds, 3) if seconds is not None else None),
            _py_scalar(round(minutes, 3) if minutes is not None else None),
            ann.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            vals['r'], vals['g'], vals['b'],
        ]
        if has_log:
            row_values.extend([vals['log_br'], vals['log_bg']])
        row_values.append(ann.description)

        for col, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill = row_fill

        color_cell = ws.cell(row=row_idx, column=color_col, value="")
        color_cell.fill = _hex_to_fill(ann.color)


def _write_markers_block(ws, annotations, saved_data=None, has_log=True, start_col=17, start_row=1):
    annotations = sorted(_normalize_annotations(annotations), key=lambda a: a.time_ms)
    if not annotations:
        return start_row

    header_fill = PatternFill(
        start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
    )
    title_cell = ws.cell(row=start_row, column=start_col, value="📌 Метки")
    title_cell.font = Font(bold=True, size=13)

    headers = _marker_table_headers(has_log)
    header_row = start_row + 1
    for col_offset, header in enumerate(headers):
        cell = ws.cell(row=header_row, column=start_col + col_offset, value=header)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = header_fill

    row_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )
    color_col_offset = len(headers) - 1

    for i, ann in enumerate(annotations, 1):
        row = header_row + i
        seconds = None
        minutes = None
        try:
            seconds = float(ann.time_ms) / 1000.0
        except Exception:
            seconds = None
        try:
            minutes = float(ann.time_ms) / 60000.0
        except Exception:
            minutes = None
        vals = _annotation_values_at_time(saved_data, ann.time_ms, has_log)
        row_values = [
            i,
            _py_scalar(round(seconds, 3) if seconds is not None else None),
            _py_scalar(round(minutes, 3) if minutes is not None else None),
            ann.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            vals['r'], vals['g'], vals['b'],
        ]
        if has_log:
            row_values.extend([vals['log_br'], vals['log_bg']])
        row_values.append(ann.description)

        for col_offset, value in enumerate(row_values):
            cell = ws.cell(row=row, column=start_col + col_offset, value=value)
            cell.fill = row_fill

        color_cell = ws.cell(row=row, column=start_col + color_col_offset, value="")
        color_cell.fill = _hex_to_fill(ann.color)

    for col_offset, width in enumerate(_marker_table_widths(has_log)):
        ws.column_dimensions[get_column_letter(start_col + col_offset)].width = width

    return header_row + len(annotations) + 2


def _annotation_row_map(saved_data, annotations):
    if not saved_data or not annotations:
        return {}
    row_to_anns = {}
    for ann in annotations:
        idx = _nearest_saved_data_index(saved_data, ann.time_ms)
        row_to_anns.setdefault(idx, []).append(ann)
    return row_to_anns


def _annotation_column_layout(has_log: bool):
    if has_log:
                                                                                      
        return {'time': 12, 'marker_rgb': 16, 'marker_log': 17}
    return {'time': 4, 'marker_rgb': 5, 'marker_log': None}


def _apply_marker_chart_columns(ws, saved_data, annotations, has_log, log_y_max):
    annotations = _normalize_annotations(annotations)
    if not annotations:
        return _annotation_column_layout(has_log)

    layout = _annotation_column_layout(has_log)
    row_map = _annotation_row_map(saved_data, annotations)

    for col, header in ((layout['marker_rgb'], 'Метка ▲'),):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(
            start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
        )

    if has_log and layout['marker_log']:
        cell = ws.cell(row=1, column=layout['marker_log'], value='Метка ▲ Log')
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(
            start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
        )

    for row_idx in row_map:
        excel_row = row_idx + 2
        ws.cell(row=excel_row, column=layout['marker_rgb'], value=RGB_Y_MAX)
        if has_log and layout['marker_log']:
            ws.cell(row=excel_row, column=layout['marker_log'], value=float(log_y_max))

    ws.column_dimensions[get_column_letter(layout['marker_rgb'])].width = 10
    if has_log and layout['marker_log']:
        ws.column_dimensions[get_column_letter(layout['marker_log'])].width = 12

    return layout


                                           

def _sheet_headers_for_markers(has_log: bool = True) -> List[str]:
    headers = [
        '#', 'Тип', 'На графике (с)', 'На графике (мин)', 'Создано',
        'R', 'G', 'B',
    ]
    if has_log:
        headers.extend(['Log10(B/R)', 'Log10(B/G)'])
    headers.extend(['RGB_sum_slope_30s', 'Transition_score', 'Описание', 'Цвет'])
    return headers


def _fill_all_marker_sheet_headers(ws, has_log: bool = True):
    headers = _sheet_headers_for_markers(has_log)
    widths = [5, 12, 15, 15, 20, 9, 9, 9]
    if has_log:
        widths.extend([13, 13])
    widths.extend([18, 17, 34, 10])
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_marker_rows(ws, annotations, saved_data=None, has_log=True, is_auto=None):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    anns = _filter_annotations(annotations, is_auto=is_auto)
    headers = _sheet_headers_for_markers(has_log)
    color_col = len(headers)
    row_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    auto_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")

    for row_idx, ann in enumerate(anns, 2):
        seconds = _py_scalar(round(float(ann.time_ms) / 1000.0, 3)) if ann.time_ms is not None else None
        minutes = _py_scalar(round(float(ann.time_ms) / 60000.0, 3)) if ann.time_ms is not None else None
        vals = _annotation_values_at_time(saved_data, ann.time_ms, has_log)
        row_values = [
            row_idx - 1,
            'Авто' if ann.is_auto else 'Ручная',
            seconds,
            minutes,
            ann.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            vals['r'] if not ann.is_auto else _py_scalar(ann.r),
            vals['g'] if not ann.is_auto else _py_scalar(ann.g),
            vals['b'] if not ann.is_auto else _py_scalar(ann.b),
        ]
        if has_log:
            row_values.extend([vals['log_br'], vals['log_bg']])
        row_values.extend([
            _py_scalar(ann.rgb_sum_slope_30s),
            _py_scalar(ann.transition_score),
            ann.description,
        ])
        for col, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill = auto_fill if ann.is_auto else row_fill
        ws.cell(row=row_idx, column=color_col, value='').fill = _hex_to_fill(ann.color)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"


def _write_graph_data_sheet(ws, saved_data, custom_graphs=None, custom_variables=None):
    custom_graphs = list(custom_graphs or [])
    custom_variables = list(custom_variables or [])
    has_log = len(saved_data) > 0 and len(saved_data[0]) > 4
    if has_log:
        headers = ['R', 'G', 'B', 'Log10(B/R)', 'Log10(B/G)', 'RGB_sum',
                   'R_frac', 'G_frac', 'B_frac',
                   'time,ms', 'RGB_sum_smooth_10s', 'RGB_sum_slope_10s', 'RGB_sum_slope_30s',
                   'RGB_vector_speed_30s', 'Chromaticity_speed_30s', 'K_chrom_from_previous_state',
                   'Log_ratio_speed_30s', 'RGB_sum_acceleration_30s', 'Transition_score']
    else:
        headers = ['R', 'G', 'B', 'time,ms']

    valid_custom = []
    for definition in custom_graphs:
        name = str(definition.get('name', '')).strip()
        formula = str(definition.get('formula', '')).strip()
        if name and formula:
            valid_custom.append((name, formula))
            headers.append(name)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    # Base arrays used by the custom formula engine. Custom formulas are evaluated
    # during export so they are persisted in Excel exactly like built-in series.
    if saved_data:
        # One compact numeric block replaces both an unused object matrix and
        # many separately-built column arrays. This lowers peak export memory
        # for long sessions while keeping column views usable by custom formulas.
        try:
            numeric_data = np.asarray(saved_data, dtype=np.float64)
            if numeric_data.ndim != 2:
                numeric_data = None
        except (TypeError, ValueError, OverflowError):
            numeric_data = None

        def col(i):
            try:
                if numeric_data is not None and numeric_data.shape[1] > i:
                    return numeric_data[:, i]
                return np.asarray([float(row[i]) if row[i] is not None else np.nan for row in saved_data], dtype=np.float64)
            except Exception:
                return np.full(len(saved_data), np.nan, dtype=np.float64)
        t_arr = col(0)
        r_arr = col(1)
        g_arr = col(2)
        b_arr = col(3)
        log_br_arr = col(4) if has_log else np.full(len(saved_data), np.nan)
        log_bg_arr = col(5) if has_log else np.full(len(saved_data), np.nan)
        rgb_sum_arr = col(6) if has_log else r_arr + g_arr + b_arr
        slope30_arr = col(14) if has_log else np.full(len(saved_data), np.nan)
        transition_arr = col(20) if has_log else np.full(len(saved_data), np.nan)

        env = {
            'np': np,
            'log10': np.log10,
            'sqrt': np.sqrt,
            'exp': np.exp,
            'abs': np.abs,
            'sin': np.sin,
            'cos': np.cos,
            'mean': np.mean,
            'diff': lambda x: np.gradient(x) if len(x) else x,
            't': t_arr, 'T': t_arr,
            'r': r_arr, 'R': r_arr,
            'g': g_arr, 'G': g_arr,
            'b': b_arr, 'B': b_arr,
            'log_br': log_br_arr,
            'log_bg': log_bg_arr,
            'rgb_sum': rgb_sum_arr, 'RGB_sum': rgb_sum_arr, 'RGB_SUM': rgb_sum_arr,
            'slope30': slope30_arr,
            'transition_score': transition_arr,
            'rgb_vector_speed30': col(15) if has_log else np.full(len(saved_data), np.nan),
            'chromaticity_speed30': col(16) if has_log else np.full(len(saved_data), np.nan),
            'k_chrom_previous': col(17) if has_log else np.full(len(saved_data), np.nan),
            'log_ratio_speed30': col(18) if has_log else np.full(len(saved_data), np.nan),
            'rgb_sum_acceleration30': col(19) if has_log else np.full(len(saved_data), np.nan),
        }
        safe_builtins = {'abs': abs, 'min': min, 'max': max, 'round': round, 'float': float, 'int': int}

        # Evaluate custom variables first
        for cv in custom_variables:
            v_name = str(cv.get('name', '')).strip()
            v_formula = str(cv.get('formula', '')).strip()
            if v_name and v_formula and v_name.isidentifier():
                try:
                    v_val = eval(v_formula, {'__builtins__': safe_builtins}, env)
                    env[v_name] = np.asarray(v_val, dtype=np.float64)
                except Exception:
                    pass

        custom_values = []
        for name, formula in valid_custom:
            try:
                value = eval(formula, {'__builtins__': safe_builtins}, env)
                if np.isscalar(value):
                    values = np.full(len(saved_data), float(value), dtype=np.float64)
                else:
                    values = np.asarray(value, dtype=np.float64).reshape(-1)
                    if len(values) != len(saved_data):
                        raise ValueError(f'формула вернула {len(values)} точек вместо {len(saved_data)}')
                custom_values.append(values)
            except Exception as exc:
                print(f"⚠️ Пользовательский график '{name}' не экспортирован: {exc}")
                custom_values.append(np.full(len(saved_data), np.nan, dtype=np.float64))
    else:
        custom_values = [np.array([], dtype=np.float64) for _ in valid_custom]

    for row_idx, row_data in enumerate(saved_data, 2):
        if has_log:
            base_values = [row_data[i] if i != 11 else row_data[0] for i in range(1, 21)]
            # Explicit order avoids accidental shifts if saved_data changes.
            base_values = [
                row_data[1], row_data[2], row_data[3], row_data[4], row_data[5], row_data[6],
                row_data[7], row_data[8], row_data[9], row_data[0],
                row_data[12], row_data[13], row_data[14], row_data[15], row_data[16], row_data[17],
                row_data[18], row_data[19], row_data[20]
            ]
        else:
            base_values = [row_data[1], row_data[2], row_data[3], row_data[0]]
        for col_idx, value in enumerate(base_values, 1):
            ws.cell(row=row_idx, column=col_idx, value=_py_scalar(value))
        for j, values in enumerate(custom_values, len(base_values) + 1):
            ws.cell(row=row_idx, column=j, value=_py_scalar(values[row_idx - 2]) if len(values) else None)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18
    ws.column_dimensions['L'].width = 12 if has_log else 18
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    return has_log


def _ensure_export_sheets(wb, has_log=True):
    desired = ['Данные графиков', 'Все метки', 'Ручные метки', 'Автометки']
    for name in desired:
        if name in wb.sheetnames:
            ws = wb[name]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = wb.create_sheet(name)
    # Remove legacy sheets from the old export format so the workbook has a clear
    # four-sheet structure after migration.
    for legacy in ('RGB Data', 'Annotations', 'метки(авт.)'):
        if legacy in wb.sheetnames and legacy not in desired:
            del wb[legacy]
    # Reorder explicitly.
    for idx, name in enumerate(desired):
        wb._sheets.insert(idx, wb._sheets.pop(wb._sheets.index(wb[name])))
    _fill_all_marker_sheet_headers(wb['Все метки'], has_log)
    _fill_all_marker_sheet_headers(wb['Ручные метки'], has_log)
    _fill_all_marker_sheet_headers(wb['Автометки'], has_log)
    return wb


def create_session_excel(save_folder: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = os.path.join(save_folder, f"analysis_{timestamp}.xlsx")
    os.makedirs(save_folder, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _ensure_export_sheets(wb, has_log=True)
    wb.save(filename)
    return filename


def save_annotations_to_workbook(excel_path: str, annotation_manager: AnnotationManager) -> Optional[str]:
    if not excel_path:
        return None
    annotations = annotation_manager.get_annotations()
    wb = load_workbook(excel_path) if os.path.isfile(excel_path) else Workbook()
    if 'RGB Data' in wb.sheetnames or 'Annotations' in wb.sheetnames or 'метки(авт.)' in wb.sheetnames:
        # Migrate an old-format workbook without carrying the marker block onto the data sheet.
        has_log = True
    else:
        has_log = 'Данные графиков' in wb.sheetnames
    _ensure_export_sheets(wb, has_log=has_log)
    _write_marker_rows(wb['Все метки'], annotations, saved_data=None, has_log=has_log, is_auto=None)
    _write_marker_rows(wb['Ручные метки'], annotations, saved_data=None, has_log=has_log, is_auto=False)
    _write_marker_rows(wb['Автометки'], annotations, saved_data=None, has_log=has_log, is_auto=True)
    wb.save(excel_path)
    return excel_path


def save_annotations_to_excel(save_folder: str, annotation_manager: AnnotationManager,
                              filename_base: str, excel_path: Optional[str] = None) -> Optional[str]:
    annotations = annotation_manager.get_annotations()
    if excel_path:
        return save_annotations_to_workbook(excel_path, annotation_manager)
    if not annotations:
        return None
    filename = os.path.join(save_folder, f"{filename_base}_annotations.xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    _ensure_export_sheets(wb, has_log=True)
    _write_marker_rows(wb['Все метки'], annotations, has_log=True, is_auto=None)
    _write_marker_rows(wb['Ручные метки'], annotations, has_log=True, is_auto=False)
    _write_marker_rows(wb['Автометки'], annotations, has_log=True, is_auto=True)
    wb.save(filename)
    return filename


def save_to_excel(save_folder: str, saved_data: List, excel_path: Optional[str] = None,
                  annotations: Optional[List] = None, custom_graphs: Optional[List[Dict[str, str]]] = None,
                  custom_variables: Optional[List[Dict[str, str]]] = None) -> str:
    if excel_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.join(save_folder, f"analysis_{timestamp}.xlsx")
        wb = Workbook()
        wb.remove(wb.active)
    else:
        filename = excel_path
        wb = load_workbook(excel_path)

    has_log = len(saved_data) > 0 and len(saved_data[0]) > 4
    _ensure_export_sheets(wb, has_log=has_log)
    _write_graph_data_sheet(wb['Данные графиков'], saved_data, custom_graphs=custom_graphs, custom_variables=custom_variables)

    annotations = _normalize_annotations(annotations)
    _write_marker_rows(wb['Все метки'], annotations, saved_data=saved_data, has_log=has_log, is_auto=None)
    _write_marker_rows(wb['Ручные метки'], annotations, saved_data=saved_data, has_log=has_log, is_auto=False)
    _write_marker_rows(wb['Автометки'], annotations, saved_data=saved_data, has_log=has_log, is_auto=True)

    os.makedirs(save_folder, exist_ok=True)
    wb.save(filename)
    if not os.path.isfile(filename):
        raise IOError(f"Файл не появился после сохранения: {filename}")
    return filename

def save_graph_to_image(save_folder: str, saved_data: List, filename_base: str) -> Optional[str]:
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
    except ImportError:
        print("matplotlib не установлен, график не сохранён")
        return None

    if not saved_data:
        return None

    times = [row[0] for row in saved_data]
    r_vals = [row[1] for row in saved_data]
    g_vals = [row[2] for row in saved_data]
    b_vals = [row[3] for row in saved_data]

    has_log = len(saved_data[0]) > 4
    log_vals = [row[4] for row in saved_data] if has_log else None

    if has_log:
        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(16, 10))

        ax1.plot(times, r_vals, color='red', linewidth=1.5, label='R', alpha=0.8)
        ax1.plot(times, g_vals, color='green', linewidth=1.5, label='G', alpha=0.8)
        ax1.plot(times, b_vals, color='blue', linewidth=1.5, label='B', alpha=0.8)

        ax1.set_xlabel('Time (ms)', fontsize=12)
        ax1.set_ylabel('Value (0-255)', fontsize=12)
        ax1.set_title('RGB Analysis - Full Graph', fontsize=14, fontweight='bold')
        ax1.legend(loc='best', fontsize=10)
        ax1.grid(True, alpha=0.3)
        ax1.set_ylim(0, 255)
        ax1.set_xlim(0, max(times) * 1.02 if times else 100)

        log_y_min, log_y_max = adaptive_log_y_range(log_vals)
        valid_indices = [i for i, v in enumerate(log_vals) if v is not None and not np.isnan(v) and not np.isinf(v)]
        if valid_indices:
            times_valid = [times[i] for i in valid_indices]
            log_vals_valid = [log_vals[i] for i in valid_indices]
            ax2.plot(times_valid, log_vals_valid, color='#CC5500', linewidth=2, label='Log10(B/R)', alpha=0.9)
        else:
            ax2.plot(times, log_vals, color='#CC5500', linewidth=2, label='Log10(B/R)', alpha=0.9)

        ax2.set_xlabel('Time (ms)', fontsize=12)
        ax2.set_ylabel('Log10(B/R)', fontsize=12)
        ax2.set_title('Log10(B/R) Analysis', fontsize=14, fontweight='bold')
        ax2.legend(loc='best', fontsize=10)
        ax2.grid(True, alpha=0.3)
        ax2.set_ylim(log_y_min, log_y_max)
        ax2.set_xlim(0, max(times) * 1.02 if times else 100)
    else:
        fig, ax1 = plt.subplots(figsize=(16, 6))
        ax1.plot(times, r_vals, color='red', linewidth=1.5, label='R', alpha=0.8)
        ax1.plot(times, g_vals, color='green', linewidth=1.5, label='G', alpha=0.8)
        ax1.plot(times, b_vals, color='blue', linewidth=1.5, label='B', alpha=0.8)
        ax1.set_xlabel('Time (ms)', fontsize=12)
        ax1.set_ylabel('Value (0-255)', fontsize=12)
        ax1.set_title('RGB Analysis - Full Graph', fontsize=14, fontweight='bold')
        ax1.legend(loc='best', fontsize=10)
        ax1.grid(True, alpha=0.3)
        ax1.set_ylim(0, 255)
        ax1.set_xlim(0, max(times) * 1.02 if times else 100)

    plt.tight_layout()
    graph_filename = os.path.join(save_folder, f"{filename_base}_graph.png")
    plt.savefig(graph_filename, dpi=150, bbox_inches='tight')
    plt.close(fig)
    print(f"📊 График сохранён: {graph_filename}")
    return graph_filename

